from __future__ import annotations

import io
import tempfile
import unittest
from pathlib import Path

from docx import Document as WordDocument
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from pptx import Presentation

from app.parsers import ParseError, ParserRegistry


class ParserRegistryTests(unittest.TestCase):
    """通过统一解析入口验证多格式资料都会转换为 Document。"""

    def setUp(self) -> None:
        """创建生产支持格式的解析器注册表。"""
        self.parsers = ParserRegistry.default()

    def test_text_and_markdown_are_normalized(self):
        """纯文本和 Markdown 应保留正文及统一来源字段。"""
        text_docs = self.parsers.parse("notes.txt", b"first\n\nsecond")
        markdown_docs = self.parsers.parse("guide.md", "# 标题\n正文".encode())

        self.assertEqual([item.content for item in text_docs], ["first\n\nsecond"])
        self.assertEqual(markdown_docs[0].metadata["section"], "标题")
        self.assertEqual(markdown_docs[0].metadata["filename"], "guide.md")

    def test_office_files_keep_source_locations(self):
        """Word、PPT 和 Excel 应分别保留章节、页码和 Sheet 定位信息。"""
        word = WordDocument()
        word.add_heading("第一章", level=1)
        word.add_paragraph("Word 正文")
        word_buffer = io.BytesIO()
        word.save(word_buffer)

        slides = Presentation()
        slide = slides.slides.add_slide(slides.slide_layouts[1])
        slide.shapes.title.text = "课程标题"
        slide.placeholders[1].text = "PPT 正文"
        ppt_buffer = io.BytesIO()
        slides.save(ppt_buffer)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "成绩"
        sheet.append(["姓名", "分数"])
        sheet.append(["小明", 95])
        excel_buffer = io.BytesIO()
        workbook.save(excel_buffer)

        word_docs = self.parsers.parse("course.docx", word_buffer.getvalue())
        ppt_docs = self.parsers.parse("slides.pptx", ppt_buffer.getvalue())
        excel_docs = self.parsers.parse("scores.xlsx", excel_buffer.getvalue())

        self.assertEqual(word_docs[0].metadata["section"], "第一章")
        self.assertIn("Word 正文", word_docs[0].content)
        self.assertEqual(ppt_docs[0].metadata["page"], 1)
        self.assertIn("PPT 正文", ppt_docs[0].content)
        self.assertEqual(excel_docs[0].metadata["sheet"], "成绩")
        self.assertIn("小明", excel_docs[0].content)

    def test_csv_is_parsed_as_a_table(self):
        """CSV 应按带表头的数据行输出，而不是拆散单元格。"""
        documents = self.parsers.parse("scores.csv", "姓名,分数\n小明,95\n".encode("utf-8"))

        self.assertEqual(len(documents), 1)
        self.assertIn("姓名: 小明", documents[0].content)
        self.assertEqual(documents[0].metadata["block_type"], "table")

    def test_csv_detects_header_after_notes_and_preserves_unnamed_value_columns(self):
        """说明行和空表头不得导致后续数据列静默丢失。"""
        source = """Camborne,,,,,,,,,,,,
Location 1627E 407N 87m amsl,,,,,,,,,,,,
yyyy,mm,tmax,,tmin,,af,,rain,,sun,,
,,degC,,degC,,days,,mm,,hours,,
2014,04,12.3,,7.4,,0,,81.4,,154.9,,Provisional
"""

        documents = self.parsers.parse("weather.csv", source.encode("utf-8"))

        content = documents[0].content
        self.assertIn("yyyy: 2014", content)
        self.assertIn("mm: 04", content)
        self.assertIn("tmax: 12.3", content)
        self.assertIn("tmin: 7.4", content)
        self.assertIn("rain: 81.4", content)
        self.assertIn("sun: 154.9", content)
        self.assertIn("column_13: Provisional", content)

    def test_pdf_keeps_page_number(self):
        """可复制文本 PDF 应提取正文并保留一基页码。"""
        writer = PdfWriter()
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 72 200 Td (PDF lesson text) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        buffer = io.BytesIO()
        writer.write(buffer)

        documents = self.parsers.parse("lesson.pdf", buffer.getvalue())

        self.assertEqual(documents[0].metadata["page"], 1)
        self.assertIn("PDF lesson text", documents[0].content)

    def test_unknown_or_empty_files_are_rejected(self):
        """未知格式和没有可检索文本的文件应返回可读错误。"""
        with self.assertRaisesRegex(ParseError, "不支持"):
            self.parsers.parse("archive.zip", b"data")
        with self.assertRaisesRegex(ParseError, "没有可提取"):
            self.parsers.parse("empty.txt", b"   ")

    def test_supported_extensions_match_v2_scope(self):
        """解析入口应公开 V2 首期承诺的全部扩展名。"""
        self.assertEqual(
            self.parsers.supported_extensions,
            {".pdf", ".docx", ".md", ".txt", ".pptx", ".xlsx", ".csv"},
        )


if __name__ == "__main__":
    unittest.main()
