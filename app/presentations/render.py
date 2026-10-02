"""有限布局的 PPTX 生成，文字容量超出时拒绝或拆页。"""
from __future__ import annotations

import os
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt


def validate_slides(slides: list[dict], allowed_sources: set[str]) -> list[dict]:
    """大纲只引用该笔记固定版本中的来源，不接受任意图片或文件路径。"""
    if not isinstance(slides, list) or not 1 <= len(slides) <= 30:
        raise ValueError('大纲需要 1～30 页。')
    clean = []
    for slide in slides:
        title, bullets, notes = slide.get('title'), slide.get('bullets'), slide.get('notes', '')
        ids = slide.get('evidence_ids', [])
        if not isinstance(title, str) or not 1 <= len(title) <= 50:
            raise ValueError('每页标题需要 1～50 字符。')
        if not isinstance(bullets, list) or not 1 <= len(bullets) <= 6:
            raise ValueError('每页需要 1～6 条要点。')
        if any(not isinstance(x, str) or not 1 <= len(x) <= 150 for x in bullets):
            raise ValueError('单条要点不能超过 150 字符。')
        if not isinstance(notes, str) or len(notes) > 5000 or not isinstance(ids, list) or any(not isinstance(x, str) or x not in allowed_sources for x in ids):
            raise ValueError('讲解备注或来源无效。')
        table = slide.get('table')
        if table is not None:
            if not isinstance(table, dict) or not isinstance(table.get('headers'), list) or not 1 <= len(table['headers']) <= 6:
                raise ValueError('表格需要 1～6 列。')
            rows = table.get('rows')
            if not isinstance(rows, list) or not 1 <= len(rows) <= 6 or any(not isinstance(row, list) or len(row) != len(table['headers']) for row in rows):
                raise ValueError('表格需要 1～6 行，行列数必须匹配。')
            if any(not isinstance(cell, str) or len(cell) > 35 for row in [table['headers'], *rows] for cell in row):
                raise ValueError('表格单元格不能超过 35 字符。')
        clean.append({'title': title, 'bullets': bullets, 'notes': notes,
                      'evidence_ids': list(dict.fromkeys(ids)), **({'table': table} if table else {})})
    return clean


def render_pptx(slides: list[dict], evidence: list[dict], target: Path) -> dict:
    """按固定字号和保守行容量拆页，原子发布并重新打开验证文件结构。"""
    sources = {s['id']: s for s in evidence}
    slides = validate_slides(slides, set(sources))
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    labels = {key: f'P{i}' for i, key in enumerate(sources, 1)}
    pages = 0
    for original in slides:
        groups, group, weight = [], [], 0
        for bullet in original['bullets']:
            cost = max(1, (len(bullet) + 31) // 32)
            if group and weight + cost > 8:
                groups.append(group)
                group, weight = [], 0
            group.append(bullet)
            weight += cost
        groups.append(group)
        for index, bullets in enumerate(groups):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            title = original['title'] + ('（续）' if index else '')
            box = slide.shapes.add_textbox(Inches(.6), Inches(.4), Inches(12.1), Inches(1.1))
            box.text_frame.text = title
            for paragraph in box.text_frame.paragraphs:
                paragraph.font.size = Pt(32)
                paragraph.font.name = 'Microsoft YaHei'
            box = slide.shapes.add_textbox(Inches(.8), Inches(1.7), Inches(11.7), Inches(4.45))
            box.text_frame.word_wrap = True
            for i, bullet in enumerate(bullets):
                paragraph = box.text_frame.paragraphs[0] if i == 0 else box.text_frame.add_paragraph()
                paragraph.text = '• ' + bullet
                paragraph.font.size = Pt(22)
                paragraph.font.name = 'Microsoft YaHei'
                paragraph.space_after = Pt(12)
            refs = []
            for eid in original['evidence_ids']:
                source = sources[eid]
                location = f"PDF p.{source['page']}" if source['source_type'] == 'paper' else source['url']
                refs.append(f"[{labels[eid]}] {location}")
            footer = slide.shapes.add_textbox(Inches(.7), Inches(6.5), Inches(11.9), Inches(.6))
            footer.text_frame.word_wrap = True
            footer.text_frame.text = '来源：' + ('；'.join(refs)[:240] if refs else '用户笔记 / 辅助解释，请对照原文核验')
            footer.text_frame.paragraphs[0].font.size = Pt(11)
            slide.notes_slide.notes_text_frame.text = original['notes'] + '\n\n' + '\n'.join(
                f"[{labels[eid]}] {sources[eid].get('filename', sources[eid].get('title', ''))}\n"
                f"{sources[eid].get('url', '')} 摘录：{sources[eid]['quote']}" for eid in original['evidence_ids'])
            pages += 1
            if pages > 60:
                raise ValueError('拆分后的 PPT 超过 60 页。')
        if original.get('table'):
            pages += add_table_pages(deck, original, sources, labels)
            if pages > 60:
                raise ValueError('拆分后的 PPT 超过 60 页。')
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = target.with_suffix('.tmp.pptx')
    try:
        deck.save(temporary)
        reopened = Presentation(temporary)
        if len(reopened.slides) != pages:
            raise ValueError('PPTX 结构验证失败。')
        for slide in reopened.slides:
            for shape in slide.shapes:
                if shape.left < 0 or shape.top < 0 or shape.left + shape.width > reopened.slide_width or shape.top + shape.height > reopened.slide_height:
                    raise ValueError('PPTX 文字框超出画布。')
        os.replace(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)
    return {'page_count': pages, 'visual_validation': 'pending',
            'limitations': ['已检查文件结构和文字框边界；需实际渲染检查换行、字体及溢出。']}


def add_table_pages(deck, original: dict, sources: dict, labels: dict) -> int:
    """按保守中文字宽计算行高，长表分多页并重复表头和来源。"""
    headers, rows = original['table']['headers'], original['table']['rows']
    capacity = max(4, int((11.9 * 72 / len(headers) - 15) / 18))
    def height(row):
        """每行预留文字行距和上下内边距，不能缩字号塞满表格。"""
        return max(36, max((len(cell) + capacity - 1) // capacity for cell in row) * 25 + 12)
    header_height = height(headers)
    groups, group, size = [], [], header_height
    for row in rows:
        if group and size + height(row) > 4.7 * 72:
            groups.append(group)
            group, size = [], header_height
        if size + height(row) > 4.7 * 72:
            raise ValueError('表头和单行内容超过画布，请缩短表格文本。')
        group.append(row)
        size += height(row)
    if group:
        groups.append(group)
    for index, group in enumerate(groups):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        heading = slide.shapes.add_textbox(Inches(.6), Inches(.4), Inches(12.1), Inches(1.1))
        heading.text_frame.text = original['title'] + ('（表格续）' if index else '（表格）')
        heading.text_frame.paragraphs[0].font.size = Pt(28)
        matrix = [headers, *group]
        shape = slide.shapes.add_table(len(matrix), len(headers), Inches(.7), Inches(1.65), Inches(11.9),
                                      Pt(sum(height(row) for row in matrix)))
        for r, row in enumerate(matrix):
            shape.table.rows[r].height = Pt(height(row))
            for c, value in enumerate(row):
                cell = shape.table.cell(r, c)
                cell.text = value
                cell.text_frame.word_wrap = True
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.size = Pt(18)
                    paragraph.font.name = 'Microsoft YaHei'
                    paragraph.font.bold = r == 0
        footer = slide.shapes.add_textbox(Inches(.7), Inches(6.5), Inches(11.9), Inches(.6))
        footer.text_frame.text = '来源：' + '；'.join(f"[{labels[eid]}] " +
             (f"PDF p.{sources[eid]['page']}" if sources[eid]['source_type'] == 'paper' else sources[eid]['url'])
             for eid in original['evidence_ids'])[:240]
        footer.text_frame.paragraphs[0].font.size = Pt(11)
        slide.notes_slide.notes_text_frame.text = original['notes'] + '\n' + '\n'.join(
            f"[{labels[eid]}] {sources[eid].get('filename', sources[eid].get('title', ''))} 摘录：{sources[eid]['quote']}"
            for eid in original['evidence_ids'])
    return len(groups)
