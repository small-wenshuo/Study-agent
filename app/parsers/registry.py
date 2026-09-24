from __future__ import annotations

import csv
import io
import re
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from docx import Document as WordDocument
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation

from app.parsers.base import Document, DocumentParser, ParseError


def _metadata(filename: str, block_type: str = "paragraph", **extra: Any) -> dict[str, Any]:
    """生成所有解析器共用的来源字段。"""
    return {
        "filename": Path(filename).name,
        "file_type": Path(filename).suffix.lower().lstrip("."),
        "page": None,
        "sheet": None,
        "section": None,
        "block_type": block_type,
        **extra,
    }


def _decode_text(data: bytes) -> str:
    """优先按 UTF-8 解码，兼容常见中文本地编码。"""
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ParseError("文本编码不受支持，请转换为 UTF-8")


class PlainTextParser:
    """解析 TXT 文本。"""

    extensions = frozenset({".txt"})

    def parse(self, filename: str, data: bytes) -> list[Document]:
        """把整个短文本作为一个结构块输出。"""
        content = _decode_text(data).strip()
        return [Document(content, _metadata(filename))] if content else []


class MarkdownParser:
    """按 Markdown 标题切分章节并保留标题定位。"""

    extensions = frozenset({".md"})
    _heading = re.compile(r"^#{1,6}\s+(.+?)\s*$")

    def parse(self, filename: str, data: bytes) -> list[Document]:
        """将每个标题区段转换为独立 Document。"""
        documents: list[Document] = []
        section: str | None = None
        lines: list[str] = []

        def flush() -> None:
            content = "\n".join(lines).strip()
            if content:
                documents.append(Document(content, _metadata(filename, section=section)))

        for line in _decode_text(data).splitlines():
            match = self._heading.match(line)
            if match:
                flush()
                lines = []
                section = match.group(1).strip()
            else:
                lines.append(line)
        flush()
        return documents


class CsvParser:
    """把 CSV 表头与数据行组合为可检索表格文本。"""

    extensions = frozenset({".csv"})

    def parse(self, filename: str, data: bytes) -> list[Document]:
        """保留表头语义，避免将单元格拆成无上下文片段。"""
        reader = csv.DictReader(io.StringIO(_decode_text(data)))
        rows = [
            " | ".join(f"{key}: {value}" for key, value in row.items() if key and value not in (None, ""))
            for row in reader
        ]
        content = "\n".join(row for row in rows if row).strip()
        return [Document(content, _metadata(filename, "table"))] if content else []


class PdfParser:
    """逐页解析可复制文本的 PDF。"""

    extensions = frozenset({".pdf"})

    def parse(self, filename: str, data: bytes) -> list[Document]:
        """为每个非空页面保留一基页码。"""
        try:
            reader = PdfReader(io.BytesIO(data))
            return [
                Document(text, _metadata(filename, page=index))
                for index, page in enumerate(reader.pages, start=1)
                if (text := (page.extract_text() or "").strip())
            ]
        except Exception as exc:
            raise ParseError("PDF 解析失败或文件已损坏") from exc


class DocxParser:
    """解析 Word 标题、段落和表格。"""

    extensions = frozenset({".docx"})

    def parse(self, filename: str, data: bytes) -> list[Document]:
        """按标题聚合正文，同时独立保留表格内容。"""
        try:
            source = WordDocument(io.BytesIO(data))
        except Exception as exc:
            raise ParseError("DOCX 解析失败或文件已损坏") from exc
        documents: list[Document] = []
        section: str | None = None
        paragraphs: list[str] = []

        def flush() -> None:
            content = "\n".join(paragraphs).strip()
            if content:
                documents.append(Document(content, _metadata(filename, section=section)))

        for paragraph in source.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            if paragraph.style and paragraph.style.name.startswith("Heading"):
                flush()
                paragraphs = []
                section = text
            else:
                paragraphs.append(text)
        flush()
        for table in source.tables:
            rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
            content = "\n".join(row for row in rows if row.strip(" |"))
            if content:
                documents.append(Document(content, _metadata(filename, "table", section=section)))
        return documents


class PptxParser:
    """逐页解析 PowerPoint 文本。"""

    extensions = frozenset({".pptx"})

    def parse(self, filename: str, data: bytes) -> list[Document]:
        """保留幻灯片页码和标题，忽略空页面。"""
        try:
            presentation = Presentation(io.BytesIO(data))
        except Exception as exc:
            raise ParseError("PPTX 解析失败或文件已损坏") from exc
        documents: list[Document] = []
        for index, slide in enumerate(presentation.slides, start=1):
            texts = [
                shape.text.strip()
                for shape in slide.shapes
                if getattr(shape, "has_text_frame", False) and shape.text.strip()
            ]
            if texts:
                title = slide.shapes.title.text.strip() if slide.shapes.title else None
                documents.append(Document("\n".join(texts), _metadata(filename, page=index, section=title)))
        return documents


class XlsxParser:
    """逐 Sheet 解析 Excel 表格。"""

    extensions = frozenset({".xlsx"})

    def parse(self, filename: str, data: bytes) -> list[Document]:
        """使用首个非空行作为表头，并保留 Sheet 名称。"""
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            raise ParseError("XLSX 解析失败或文件已损坏") from exc
        documents: list[Document] = []
        for sheet in workbook.worksheets:
            values = [list(row) for row in sheet.iter_rows(values_only=True) if any(cell is not None for cell in row)]
            if len(values) < 2:
                continue
            headers = [str(cell) if cell is not None else f"列{index}" for index, cell in enumerate(values[0], start=1)]
            rows = [
                " | ".join(
                    f"{headers[index]}: {cell}"
                    for index, cell in enumerate(row)
                    if index < len(headers) and cell not in (None, "")
                )
                for row in values[1:]
            ]
            content = "\n".join(row for row in rows if row)
            if content:
                documents.append(Document(content, _metadata(filename, "table", sheet=sheet.title)))
        workbook.close()
        return documents


class ParserRegistry:
    """按扩展名分派解析器并统一处理空结果。"""

    def __init__(self, parsers: Iterable[DocumentParser]) -> None:
        """登记解析器，拒绝重复扩展名。"""
        self._parsers: dict[str, DocumentParser] = {}
        for parser in parsers:
            for extension in parser.extensions:
                if extension in self._parsers:
                    raise ValueError(f"重复解析器扩展名：{extension}")
                self._parsers[extension] = parser

    @classmethod
    def default(cls) -> "ParserRegistry":
        """构造 V2 首期全部格式的生产注册表。"""
        return cls([PlainTextParser(), MarkdownParser(), CsvParser(), PdfParser(), DocxParser(), PptxParser(), XlsxParser()])

    @property
    def supported_extensions(self) -> set[str]:
        """返回可用于上传校验的扩展名集合。"""
        return set(self._parsers)

    def parse(self, filename: str, data: bytes) -> list[Document]:
        """分派格式解析器，并拒绝空文件或无正文结果。"""
        if not data:
            raise ParseError("文件为空")
        extension = Path(filename).suffix.lower()
        parser = self._parsers.get(extension)
        if parser is None:
            raise ParseError(f"不支持的文件格式：{extension or '无扩展名'}")
        documents = parser.parse(Path(filename).name, data)
        if not documents:
            raise ParseError("文件中没有可提取的文本；扫描版 PDF 暂不支持")
        return documents
